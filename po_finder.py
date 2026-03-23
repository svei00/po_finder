import os
import shutil
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import threading
import xml.etree.ElementTree as ET
import pandas as pd
from collections import defaultdict
from datetime import datetime

# ── Namespaces common in Mexican CFDI XMLs ───────────────────────────────────
CFDI_NS  = "http://www.sat.gob.mx/cfd/4"
CFDI3_NS = "http://www.sat.gob.mx/cfd/3"

# ── Colour palette ────────────────────────────────────────────────────────────
BG      = "#1e1e2e"
CARD    = "#2a2a3e"
ACCENT  = "#7c6af7"
ACCENT2 = "#5dd6b5"
YELLOW  = "#f9e04b"
TEXT    = "#e0e0f0"
SUBTEXT = "#8888aa"
RED     = "#f28b82"
BTN_FG  = "#ffffff"

HEADERS = [
    "Fecha", "PO", "Serie", "Folio",
    "UUID", "UUID Relacionado",
    "RFC Receptor", "Nombre / Denominacion Social",
    "Subtotal", "Descuento", "IVA", "Total",
    "Concepto"
]

# Timbre Fiscal Digital namespace
TFD_NS = "http://www.sat.gob.mx/TimbreFiscalDigital"


# ─────────────────────────────────────────────────────────────────────────────
def read_po_numbers(filepath, column):
    ext = os.path.splitext(filepath)[1].lower()
    if ext in (".xls", ".xlsx"):
        df = pd.read_excel(filepath, dtype=str)
    elif ext == ".csv":
        df = pd.read_csv(filepath, dtype=str)
    else:
        try:
            df = pd.read_csv(filepath, sep="\t", dtype=str)
        except Exception:
            df = pd.read_csv(filepath, dtype=str)
    return df[column].dropna().str.strip().unique().tolist()


def _find_node(root, tag):
    """Find a CFDI node by tag, trying v4 then v3 namespace."""
    for ns in [CFDI_NS, CFDI3_NS]:
        node = root.find(f"{{{ns}}}{tag}")
        if node is None:
            node = root.find(f".//{{{ns}}}{tag}")
        if node is not None:
            return node, ns
    return None, None


def _find_all_nodes(root, tag):
    """Find all CFDI nodes by tag, trying v4 then v3 namespace."""
    for ns in [CFDI_NS, CFDI3_NS]:
        nodes = root.findall(f".//{{{ns}}}{tag}")
        if nodes:
            return nodes, ns
    return [], None


def _get(node, *attrs):
    """Return first non-None attribute value from a node."""
    if node is None:
        return ""
    for a in attrs:
        v = node.get(a)
        if v is not None:
            return v.strip()
    return ""


def extract_invoice_data(xml_path, matched_pos):
    """
    Parse a CFDI XML and return a dict with all report fields.
    matched_pos: list of PO strings found in this file.
    """
    row = {h: "" for h in HEADERS}
    row["PO"] = " | ".join(sorted(matched_pos))

    try:
        tree = ET.parse(xml_path)
        root = tree.getroot()

        # ── Comprobante root attributes ───────────────────────────────────────
        # The root IS the Comprobante node in CFDI
        comp = root
        # Try to detect namespace from root tag
        ns = CFDI_NS
        if CFDI3_NS in root.tag:
            ns = CFDI3_NS

        row["Fecha"]    = _get(comp, "Fecha")
        row["Serie"]    = _get(comp, "Serie")
        row["Folio"]    = _get(comp, "Folio")
        row["Subtotal"] = _get(comp, "SubTotal", "Subtotal")
        row["Total"]    = _get(comp, "Total")

        descuento = _get(comp, "Descuento")
        row["Descuento"] = descuento if descuento else "0.00"

        # ── Receptor ──────────────────────────────────────────────────────────
        receptor, _ = _find_node(root, "Receptor")
        if receptor is not None:
            row["RFC Receptor"]                  = _get(receptor, "Rfc", "rfc")
            row["Nombre / Denominacion Social"]  = _get(receptor, "Nombre", "nombre")

        # ── IVA – Traslado Impuesto 002 ───────────────────────────────────────
        traslados, _ = _find_all_nodes(root, "Traslado")
        iva_total = 0.0
        for t in traslados:
            imp = _get(t, "Impuesto")
            if imp == "002":
                importe = _get(t, "Importe")
                try:
                    iva_total += float(importe)
                except ValueError:
                    pass
        row["IVA"] = f"{iva_total:.2f}" if iva_total else "0.00"

        # ── UUID from Timbre Fiscal Digital ──────────────────────────────────
        tfd = root.find(f"{{{TFD_NS}}}TimbreFiscalDigital")
        if tfd is None:
            tfd = root.find(f".//{{{TFD_NS}}}TimbreFiscalDigital")
        row["UUID"] = _get(tfd, "UUID", "uuid") if tfd is not None else ""

        # ── UUID Relacionado ──────────────────────────────────────────────────
        uuid_rels = []
        for ns in [CFDI_NS, CFDI3_NS]:
            relacionados = root.findall(f".//{{{ns}}}CfdiRelacionado")
            for rel in relacionados:
                u = _get(rel, "UUID", "uuid")
                if u:
                    uuid_rels.append(u)
            if uuid_rels:
                break
        row["UUID Relacionado"] = " | ".join(uuid_rels)

        # ── Conceptos – join all descriptions ─────────────────────────────────
        conceptos, _ = _find_all_nodes(root, "Concepto")
        descrips = []
        for c in conceptos:
            d = _get(c, "Descripcion", "descripcion")
            if d:
                descrips.append(d)
        row["Concepto"] = " | ".join(descrips)

    except Exception as e:
        row["Concepto"] = f"[XML parse error: {e}]"

    return row


def search_xml_for_po(xml_path, po_numbers):
    """
    Returns (matched_pos, emisor_rfc, emisor_nombre, receptor_rfc, receptor_nombre).
    matched_pos is a LIST of all POs found inside this XML.
    """
    try:
        tree = ET.parse(xml_path)
        root = tree.getroot()
        text_content = ET.tostring(root, encoding="unicode")

        matched_pos = [po for po in po_numbers if po in text_content]
        if not matched_pos:
            return [], None, None, None, None

        emisor,   _ = _find_node(root, "Emisor")
        receptor, _ = _find_node(root, "Receptor")

        e_rfc  = _get(emisor,   "Rfc")
        e_name = _get(emisor,   "Nombre")
        r_rfc  = _get(receptor, "Rfc")
        r_name = _get(receptor, "Nombre")

        return matched_pos, e_rfc or None, e_name or None, r_rfc or None, r_name or None

    except Exception:
        return [], None, None, None, None


def generate_xlsx(rows, xlsx_path):
    """Write the invoice report rows to an XLSX file with formatting."""
    from openpyxl import Workbook
    from openpyxl.styles import (Font, PatternFill, Alignment,
                                  Border, Side, numbers)
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"

    # ── Style definitions ─────────────────────────────────────────────────────
    hdr_font    = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    hdr_fill    = PatternFill("solid", fgColor="7C6AF7")        # ACCENT purple
    hdr_align   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    even_fill   = PatternFill("solid", fgColor="2A2A3E")
    odd_fill    = PatternFill("solid", fgColor="1E1E2E")
    cell_font   = Font(name="Arial", size=9, color="E0E0F0")
    cell_align  = Alignment(vertical="top", wrap_text=True)
    border_side = Side(style="thin", color="3A3A5E")
    thin_border = Border(bottom=border_side)

    money_fmt   = '#,##0.00'
    date_fmt    = 'YYYY-MM-DD HH:MM:SS'

    # Column widths (chars)
    col_widths = [20, 18, 8, 10, 38, 38, 16, 36, 14, 12, 14, 14, 60]

    # ── Header row ────────────────────────────────────────────────────────────
    ws.row_dimensions[1].height = 28
    for col_idx, (header, width) in enumerate(zip(HEADERS, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    # ── Data rows ─────────────────────────────────────────────────────────────
    money_cols = {"Subtotal", "Descuento", "IVA", "Total"}

    for row_idx, row_data in enumerate(rows, start=2):
        fill = odd_fill if row_idx % 2 == 0 else even_fill
        for col_idx, header in enumerate(HEADERS, start=1):
            raw = row_data.get(header, "")
            # Convert numeric strings for money columns
            if header in money_cols and raw:
                try:
                    value = float(raw)
                except ValueError:
                    value = raw
            else:
                value = raw

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = cell_font
            cell.fill      = fill
            cell.alignment = cell_align
            cell.border    = thin_border
            if header in money_cols and isinstance(value, float):
                cell.number_format = money_fmt

        ws.row_dimensions[row_idx].height = 40

    # ── Totals row ────────────────────────────────────────────────────────────
    total_row = len(rows) + 2
    total_fill = PatternFill("solid", fgColor="5DD6B5")
    total_font = Font(name="Arial", bold=True, size=9, color="000000")
    total_align = Alignment(horizontal="right", vertical="center")

    ws.cell(row=total_row, column=1, value="TOTAL").font  = total_font
    ws.cell(row=total_row, column=1).fill   = total_fill
    ws.cell(row=total_row, column=1).alignment = total_align

    for col_idx, header in enumerate(HEADERS, start=1):
        if header in money_cols:
            col_letter = get_column_letter(col_idx)
            formula_cell = ws.cell(
                row=total_row, column=col_idx,
                value=f"=SUM({col_letter}2:{col_letter}{total_row - 1})"
            )
            formula_cell.number_format = money_fmt
            formula_cell.font          = total_font
            formula_cell.fill          = total_fill
            formula_cell.alignment     = Alignment(horizontal="right", vertical="center")

    ws.row_dimensions[total_row].height = 22

    wb.save(xlsx_path)


# ═══════════════════════════════════════════════════════════════════════════════
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PO -> XML Finder  |  CFDI Mexico")
        self.geometry("820x740")
        self.configure(bg=BG)
        self.resizable(True, True)

        self.data_file         = tk.StringVar()
        self.xml_dir           = tk.StringVar()
        self.dest_dir          = tk.StringVar()
        self.po_column         = tk.StringVar()
        self.copy_mode         = tk.StringVar(value="flat")
        self.available_columns = []

        self._build_ui()

    # ── UI ────────────────────────────────────────────────────────────────────
    def _build_ui(self):
        self._header()
        frame = tk.Frame(self, bg=BG, padx=24, pady=8)
        frame.pack(fill="both", expand=True)

        self._file_row(frame, "Data file (XLSX / XLS / CSV / TXT)",
                       self.data_file, self._browse_data, row=0)
        self._column_selector(frame, row=1)
        self._file_row(frame, "XML source folder  (subfolders scanned automatically)",
                       self.xml_dir, self._browse_xml_dir, row=2)
        self._file_row(frame, "Destination folder",
                       self.dest_dir, self._browse_dest_dir, row=3)
        self._copy_mode_selector(frame, row=4)
        self._run_button(frame)
        self._log_area(frame)

    def _header(self):
        hdr = tk.Frame(self, bg=ACCENT, pady=14)
        hdr.pack(fill="x")
        tk.Label(hdr, text="PO -> XML Finder",
                 font=("Segoe UI", 18, "bold"), bg=ACCENT, fg=BTN_FG).pack()
        tk.Label(hdr, text="Locate, match and copy CFDI XML files by Purchase Order number  |  Generates XLSX report",
                 font=("Segoe UI", 10), bg=ACCENT, fg="#ddddff").pack()

    def _file_row(self, parent, label, var, cmd, row):
        tk.Label(parent, text=label, bg=BG, fg=SUBTEXT,
                 font=("Segoe UI", 9)).grid(
                     row=row*3, column=0, columnspan=3, sticky="w", pady=(12, 0))
        tk.Entry(parent, textvariable=var, bg=CARD, fg=TEXT,
                 insertbackground=TEXT, relief="flat",
                 font=("Segoe UI", 10), bd=6).grid(
                     row=row*3+1, column=0, columnspan=2, sticky="ew", ipady=4)
        tk.Button(parent, text="Browse", command=cmd, bg=ACCENT, fg=BTN_FG,
                  relief="flat", font=("Segoe UI", 9, "bold"),
                  padx=12, cursor="hand2").grid(
                      row=row*3+1, column=2, sticky="ew", padx=(6, 0))
        parent.columnconfigure(0, weight=1)

    def _column_selector(self, parent, row):
        tk.Label(parent, text="PO Number column", bg=BG, fg=SUBTEXT,
                 font=("Segoe UI", 9)).grid(
                     row=row*3, column=0, columnspan=3, sticky="w", pady=(12, 0))
        inner = tk.Frame(parent, bg=BG)
        inner.grid(row=row*3+1, column=0, columnspan=3, sticky="ew")
        inner.columnconfigure(0, weight=1)

        self.col_combo = ttk.Combobox(inner, textvariable=self.po_column,
                                      state="readonly", font=("Segoe UI", 10))
        self.col_combo.grid(row=0, column=0, sticky="ew", ipady=3)
        ttk.Style().configure("TCombobox", fieldbackground=CARD,
                               background=CARD, foreground=TEXT)
        tk.Button(inner, text="Auto-detect", command=self._auto_detect_column,
                  bg=ACCENT2, fg="#000000", relief="flat",
                  font=("Segoe UI", 9, "bold"), padx=10, cursor="hand2"
                  ).grid(row=0, column=1, padx=(6, 0))

    def _copy_mode_selector(self, parent, row):
        tk.Label(parent, text="Destination structure", bg=BG, fg=SUBTEXT,
                 font=("Segoe UI", 9)).grid(
                     row=row*3, column=0, columnspan=3, sticky="w", pady=(12, 0))
        inner = tk.Frame(parent, bg=BG)
        inner.grid(row=row*3+1, column=0, columnspan=3, sticky="w")
        for text, val in [("Flat - all XMLs in one folder", "flat"),
                          ("Keep original folder structure",  "structure")]:
            tk.Radiobutton(inner, text=text, variable=self.copy_mode, value=val,
                           bg=BG, fg=TEXT, selectcolor=CARD,
                           activebackground=BG, activeforeground=ACCENT2,
                           font=("Segoe UI", 10)).pack(side="left", padx=(0, 24))

    def _run_button(self, parent):
        tk.Button(parent, text="Run - Find, Copy XMLs & Generate XLSX Report",
                  command=self._run_threaded, bg=ACCENT, fg=BTN_FG,
                  relief="flat", font=("Segoe UI", 12, "bold"),
                  pady=10, cursor="hand2").grid(
                      row=14, column=0, columnspan=3, sticky="ew", pady=(20, 6))

    def _log_area(self, parent):
        hdr = tk.Frame(parent, bg=BG)
        hdr.grid(row=15, column=0, columnspan=3, sticky="ew", pady=(4, 0))
        tk.Label(hdr, text="Log", bg=BG, fg=SUBTEXT,
                 font=("Segoe UI", 9)).pack(side="left")
        tk.Button(hdr, text="Copy Log", command=self._copy_log,
                  bg=CARD, fg=ACCENT2, relief="flat",
                  font=("Segoe UI", 8, "bold"), padx=8, cursor="hand2"
                  ).pack(side="right")
        tk.Button(hdr, text="Clear", command=self._clear_log,
                  bg=CARD, fg=SUBTEXT, relief="flat",
                  font=("Segoe UI", 8), padx=8, cursor="hand2"
                  ).pack(side="right", padx=(0, 4))

        self.progress = ttk.Progressbar(parent, mode="determinate")
        self.progress.grid(row=16, column=0, columnspan=3, sticky="ew", pady=(4, 6))

        self.log = tk.Text(parent, bg=CARD, fg=TEXT, relief="flat",
                           font=("Consolas", 9), height=12, wrap="word",
                           insertbackground=TEXT, cursor="arrow")
        self.log.grid(row=17, column=0, columnspan=3, sticky="nsew")
        sb = tk.Scrollbar(parent, command=self.log.yview, bg=CARD)
        sb.grid(row=17, column=3, sticky="ns")
        self.log["yscrollcommand"] = sb.set
        parent.rowconfigure(17, weight=1)
        self.log.bind("<Key>", self._block_typing)

    # ── Browse ────────────────────────────────────────────────────────────────
    def _browse_data(self):
        path = filedialog.askopenfilename(
            title="Select data file",
            filetypes=[("Spreadsheet / Text", "*.xlsx *.xls *.csv *.txt"),
                       ("All files", "*.*")])
        if path:
            self.data_file.set(path)
            self._load_columns(path)

    def _browse_xml_dir(self):
        d = filedialog.askdirectory(title="Select XML source folder")
        if d:
            self.xml_dir.set(d)

    def _browse_dest_dir(self):
        d = filedialog.askdirectory(title="Select destination folder")
        if d:
            self.dest_dir.set(d)

    # ── Column helpers ────────────────────────────────────────────────────────
    def _load_columns(self, path):
        try:
            ext = os.path.splitext(path)[1].lower()
            if ext in (".xls", ".xlsx"):
                df = pd.read_excel(path, nrows=0)
            elif ext == ".csv":
                df = pd.read_csv(path, nrows=0)
            else:
                try:    df = pd.read_csv(path, sep="\t", nrows=0)
                except: df = pd.read_csv(path, nrows=0)
            self.available_columns = df.columns.tolist()
            self.col_combo["values"] = self.available_columns
            for c in self.available_columns:
                if "po" in c.lower() or "order" in c.lower() or "purchase" in c.lower():
                    self.po_column.set(c)
                    break
            else:
                if self.available_columns:
                    self.po_column.set(self.available_columns[0])
        except Exception as e:
            self._log(f"[WARN] Could not read columns: {e}", RED)

    def _auto_detect_column(self):
        if not self.data_file.get():
            messagebox.showwarning("No file", "Please select a data file first.")
            return
        self._load_columns(self.data_file.get())
        self._log(f"Columns loaded: {self.available_columns}", ACCENT2)
        messagebox.showinfo("Columns detected",
                            f"Found {len(self.available_columns)} columns.\n"
                            f"Selected: [{self.po_column.get()}]\n\n"
                            "You can change it in the dropdown.")

    # ── Log helpers ───────────────────────────────────────────────────────────
    def _log(self, msg, color=None):
        if color is None:
            color = TEXT
        def _do():
            self.log.insert("end", msg + "\n", color)
            self.log.tag_configure(color, foreground=color)
            self.log.see("end")
        self.after(0, _do)

    def _copy_log(self):
        content = self.log.get("1.0", "end")
        self.clipboard_clear()
        self.clipboard_append(content)
        messagebox.showinfo("Copied", "Log copied to clipboard.")

    def _clear_log(self):
        self.log.delete("1.0", "end")
        self.progress["value"] = 0

    def _block_typing(self, event):
        allowed_keys = (
            "Up", "Down", "Left", "Right", "Home", "End",
            "Prior", "Next", "Shift_L", "Shift_R",
            "Control_L", "Control_R", "Alt_L", "Alt_R"
        )
        if event.keysym in allowed_keys:
            return
        if event.state & 0x4 and event.keysym.lower() in ("c", "a"):
            return
        return "break"

    # ── Main logic ────────────────────────────────────────────────────────────
    def _run_threaded(self):
        threading.Thread(target=self._run, daemon=True).start()

    def _run(self):
        if not self.data_file.get():
            return messagebox.showerror("Missing", "Please select a data file.")
        if not self.po_column.get():
            return messagebox.showerror("Missing", "Please select the PO column.")
        if not self.xml_dir.get():
            return messagebox.showerror("Missing", "Please select the XML source folder.")
        if not self.dest_dir.get():
            return messagebox.showerror("Missing", "Please select the destination folder.")

        # ── Check XLSX overwrite BEFORE starting ──────────────────────────────
        xlsx_path = os.path.join(self.dest_dir.get(), "reporte_facturas.xlsx")
        if os.path.exists(xlsx_path):
            answer = messagebox.askyesno(
                "XLSX already exists",
                "reporte_facturas.xlsx already exists in the destination folder.\n\n"
                "Do you want to overwrite it?\n\n"
                "YES = Overwrite\nNO  = Keep existing (report will not be generated)"
            )
            if not answer:
                self._log("XLSX report generation skipped by user.", YELLOW)
                generate_report = False
            else:
                generate_report = True
        else:
            generate_report = True

        self._log("=" * 60)
        self._log("Starting process...")

        # 1. Read PO numbers
        try:
            po_list = read_po_numbers(self.data_file.get(), self.po_column.get())
            po_set  = set(po_list)
            self._log(f"OK  PO numbers loaded: {len(po_set)}", ACCENT2)
        except Exception as e:
            self._log(f"ERROR reading data file: {e}", RED)
            return

        # 2. Walk all subfolders
        xml_files = []
        for dirpath, _, filenames in os.walk(self.xml_dir.get()):
            for f in filenames:
                if f.lower().endswith(".xml"):
                    xml_files.append(os.path.join(dirpath, f))

        if not xml_files:
            self._log("ERROR: No XML files found in folder or subfolders.", RED)
            return
        self._log(f"OK  XML files found (all subfolders): {len(xml_files)}", ACCENT2)

        os.makedirs(self.dest_dir.get(), exist_ok=True)
        self.progress["maximum"] = len(xml_files)
        self.progress["value"]   = 0

        copied        = 0
        skipped       = 0
        not_found_pos = set(po_set)
        emisor_map    = defaultdict(set)
        receptor_map  = defaultdict(set)
        xml_root      = self.xml_dir.get()
        report_rows   = []   # ← collects data for XLSX

        for i, full_path in enumerate(xml_files):
            fname = os.path.basename(full_path)
            matched_pos, e_rfc, e_name, r_rfc, r_name = search_xml_for_po(
                full_path, po_set)

            if matched_pos:
                # Determine destination path
                if self.copy_mode.get() == "structure":
                    rel       = os.path.relpath(full_path, xml_root)
                    dest_path = os.path.join(self.dest_dir.get(), rel)
                    os.makedirs(os.path.dirname(dest_path), exist_ok=True)
                else:
                    dest_path = os.path.join(self.dest_dir.get(), fname)

                rel_display = os.path.relpath(full_path, xml_root)
                po_display  = ", ".join(sorted(matched_pos))

                # Skip if already in destination
                if os.path.exists(dest_path):
                    skipped += 1
                    self._log(
                        f"  SKIP  {rel_display}  (already in destination)", YELLOW)
                else:
                    shutil.copy2(full_path, dest_path)
                    copied += 1
                    self._log(
                        f"  COPY  {rel_display}  |  PO(s): {po_display}", ACCENT2)

                # Mark ALL matched POs as found
                for po in matched_pos:
                    not_found_pos.discard(po)

                if e_rfc:
                    emisor_map[e_rfc].add(e_name or "---")
                if r_rfc:
                    receptor_map[r_rfc].add(r_name or "---")

                # ── Collect data for XLSX (always, even if file was skipped) ──
                if generate_report:
                    row_data = extract_invoice_data(full_path, matched_pos)
                    report_rows.append(row_data)

            self.progress["value"] = i + 1
            self.update_idletasks()

        # ── Generate XLSX ─────────────────────────────────────────────────────
        xlsx_msg = ""
        if generate_report and report_rows:
            try:
                generate_xlsx(report_rows, xlsx_path)
                self._log(f"XLSX saved: {xlsx_path}", ACCENT2)
                xlsx_msg = f"\n\nXLSX Report: reporte_facturas.xlsx\n({len(report_rows)} invoices)"
            except Exception as e:
                self._log(f"ERROR generating XLSX: {e}", RED)
                xlsx_msg = f"\n\nXLSX ERROR: {e}"
        elif generate_report and not report_rows:
            self._log("No matched invoices - XLSX not generated.", YELLOW)

        # ── Summary ──────────────────────────────────────────────────────────
        self._log("=" * 60)
        self._log("PROCESS COMPLETE")
        self._log(f"  Files copied  : {copied}")
        self._log(f"  Files skipped : {skipped}  (already existed in destination)")
        self._log(f"  POs matched   : {len(po_set) - len(not_found_pos)} / {len(po_set)}")

        if not_found_pos:
            self._log(
                f"  POs NOT found : {', '.join(sorted(not_found_pos))}", RED)

        emisor_lines   = []
        receptor_lines = []

        self._log("  -- Emisor / Proveedor --", ACCENT2)
        for rfc, names in emisor_map.items():
            line = f"    RFC: {rfc}  |  Nombre: {', '.join(sorted(names))}"
            emisor_lines.append(line)
            self._log(line, ACCENT2)

        self._log("  -- Receptor / Cliente --", ACCENT2)
        for rfc, names in receptor_map.items():
            line = f"    RFC: {rfc}  |  Nombre: {', '.join(sorted(names))}"
            receptor_lines.append(line)
            self._log(line, ACCENT2)

        # ── Popup ─────────────────────────────────────────────────────────────
        e_block  = "\n".join(emisor_lines)   or "  Not found in XMLs"
        r_block  = "\n".join(receptor_lines) or "  Not found in XMLs"
        nf_block = (
            f"\n\nPOs NOT found ({len(not_found_pos)}):\n"
            + "\n".join(sorted(not_found_pos))
        ) if not_found_pos else ""

        messagebox.showinfo(
            "Process Complete",
            f"Files copied : {copied}\n"
            f"Files skipped: {skipped}  (already existed)\n"
            f"POs matched  : {len(po_set) - len(not_found_pos)} / {len(po_set)}"
            f"{xlsx_msg}\n"
            f"\n-- Emisor / Proveedor --\n{e_block}"
            f"\n\n-- Receptor / Cliente --\n{r_block}"
            f"{nf_block}"
        )


if __name__ == "__main__":
    app = App()
    app.mainloop()